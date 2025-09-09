# This script is for advanced Excel processing using pandas and openpyxl
# It will be imported and used in app.py for the process_excel function
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.datavalidation import DataValidation
import datetime
from langdetect import detect, LangDetectException
import io

DEBUG_MODE = True  # Set to False to disable debug prints

def debug_print(message, extra_info=""):
    if DEBUG_MODE:
        timestamp = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        print(f"DEBUG [{timestamp}] {message}{extra_info}")

def advanced_process_excel_memory(input_data, sf_owner_id=None):
    """
    Process Excel file entirely in memory without touching disk
    """
    # Read from bytes data
    df = pd.read_excel(io.BytesIO(input_data), skiprows=4)
    debug_print("File read and rows skipped", f" (Rows: {df.shape[0]})")

    # Drop empty columns
    df = df.dropna(axis=1, how='all')
    debug_print("Empty columns dropped", f" (Columns: {df.shape[1]})")

    # Define the strict output column order (as required in output)
    strict_order = [
        'Name',
        'Buyer_Account__c',
        'CSM_Name__c',
        'Project_Manager__c',
        'Project_Manager_Email__c',
        'Date_Flagged__c',
        'Question_Visibility__c',
        'Custom_Create_Date__c',
        'Lucid_Action__c',
        'OwnerId',
        'Survey_Number__c',
        'Question_Type__c',
        'Country_Language__c',
        'QuestionID__c',
        'Required_Action__c',
        'Setup_Issue__c',
        'Recommendation__c',
        'QuestionName__c',
        'Custom_Flagged__c',
        '[D]ANSPRECODE',
        '[D]OPTION TEXT'
    ]

    # Rename input columns to match strict output names if needed
    rename_map = {
        'ANSWER PRECODE': '[D]ANSPRECODE',
        'ANSWER OPTION TEXT': '[D]OPTION TEXT'
    }
    df = df.rename(columns=rename_map)
    # Only keep columns that are in the strict_order, drop all others
    df = df[[col for col in df.columns if col in strict_order or col in rename_map.values()]]
    # Add any missing columns as empty
    for col in strict_order:
        if col not in df.columns:
            df[col] = ''
    # Reindex to strict order
    df = df[strict_order]
    debug_print("Columns aligned to strict output order", f" (Rows: {df.shape[0]}, Columns: {df.shape[1]})")

    # Fill DATE FLAGGED with current system date for all rows
    now_str = datetime.datetime.now().strftime('%Y-%m-%dT%H:%M:%S')
    df['Date_Flagged__c'] = now_str
    debug_print("Filled DATE FLAGGED with current system date for all rows", f" (Rows: {df.shape[0]})")

    # Fill Lucid_Action__c with 'Not Paused'
    df['Lucid_Action__c'] = 'Not Paused'
    
    # Fill OwnerId column with SF OwnerId if provided, else leave blank
    if 'OwnerId' in df.columns:
        df['OwnerId'] = sf_owner_id if sf_owner_id else ''

    # Detect language for each cell in 'Custom_Flagged__c' column and add a new column 'QUESTION TEXT LANGUAGE'
    if 'Custom_Flagged__c' in df.columns:
        def detect_lang_safe(text):
            try:
                if pd.notnull(text) and str(text).strip() and isinstance(text, str):
                    return detect(str(text))
                return ''
            except (LangDetectException, TypeError):
                return ''
        df['QUESTION TEXT LANGUAGE'] = df['Custom_Flagged__c'].apply(detect_lang_safe)
    debug_print("Language detection completed", f" (Rows: {df.shape[0]})")

    # Set Feedback/Recommendation/Setup_Issue__c based on language logic - always override
    if all(col in df.columns for col in ['Country_Language__c', 'QUESTION TEXT LANGUAGE', 'Required_Action__c', 'Recommendation__c', 'Setup_Issue__c']):
        mask_non_english = ~df['Country_Language__c'].astype(str).str.startswith('English')
        mask_qtext_english = df['QUESTION TEXT LANGUAGE'] == 'en'
        # Where COUNTRY LANGUAGE doesn't start with English and QUESTION TEXT is English
        mask_translate = mask_non_english & mask_qtext_english
        df.loc[mask_translate, 'Required_Action__c'] = 'Language Translation Required'
        df.loc[mask_translate, 'Recommendation__c'] = 'Translate into correct Language'
        df.loc[mask_translate, 'Setup_Issue__c'] = 'Untranslated Text'
        # Where COUNTRY LANGUAGE doesn't start with English and QUESTION TEXT is not English
        mask_ok = mask_non_english & ~mask_qtext_english
        df.loc[mask_ok, 'Required_Action__c'] = '--None--'
        df.loc[mask_ok, 'Recommendation__c'] = '--None--'
        df.loc[mask_ok, 'Setup_Issue__c'] = '--None--'
    debug_print("For non-English questions, Required_Action, Recommendation, Setup_Issue__c updated", f" (Rows: {df.shape[0]})")

    # Sort by keys if present (after language logic)
    sort_keys = [col for col in ['Required_Action__c','QuestionID__c', 'Survey_Number__c', '[D]ANSPRECODE'] if col in df.columns]
    if sort_keys:
        df = df.sort_values(by=sort_keys, kind='stable')
    debug_print("Rows sorted to bring in order by QID.", f" (Rows: {df.shape[0]})")

    # Drop 'QUESTION TEXT LANGUAGE' column before exporting
    if 'QUESTION TEXT LANGUAGE' in df.columns:
        df = df.drop(columns=['QUESTION TEXT LANGUAGE'])
    # Save to memory buffer instead of file
    output_buffer = io.BytesIO()
    df.to_excel(output_buffer, index=False, engine='openpyxl')
    debug_print("Data saved to memory buffer as Excel")
    
    # Post-process with openpyxl in memory
    output_buffer.seek(0)
    wb = openpyxl.load_workbook(output_buffer)
    ws = wb.active
    debug_print("Workbook loaded")
    # Disable word wrap, unmerge, set width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=False)
            # Reset any existing fill patterns that might interfere with sorting
            cell.fill = openpyxl.styles.PatternFill()
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))

    debug_print("Formatting applied (Disabled wordwrap, cells unmerged)")
    # Clear any existing conditional formatting that might interfere
    while ws.conditional_formatting:
        ws.conditional_formatting._cf_rules.clear()

    # debug_print(11, "Custom column widths set")
    # Set custom column widths (supports output headers directly)
    orig_col_widths = {
        'Name': 20, 'Buyer_Account__c': 20, 'CSM_Name__c': 20, 'Project_Manager__c': 20, 'Project_Manager_Email__c': 20,
        'Date_Flagged__c': 15, 'Question_Visibility__c': 10, 'Custom_Create_Date__c': 15, 'Lucid_Action__c': 20, 'OwnerId': 20,
        'Survey_Number__c': 12, 'Question_Type__c': 15, 'Country_Language__c': 17.86, 'QuestionID__c': 8.57,
        'Required_Action__c': 28.57, 'Setup_Issue__c': 28.57, 'Recommendation__c': 40, 'QuestionName__c': 20,
        'Custom_Flagged__c': 61.33, '[D]ANSPRECODE': 5.71, '[D]OPTION TEXT': 20
    }
    for cell in ws[1]:
        header = cell.value
        col_letter = get_column_letter(cell.column)
        if header in orig_col_widths:
            ws.column_dimensions[col_letter].width = orig_col_widths[header]
        else:
            ws.column_dimensions[col_letter].width = 20

    debug_print("Column widths set")
    # Format DATE CREATED (handle renamed header)
    # If the original 'DATE CREATED' was renamed to 'Custom_Create_Date__c', use df to find column index
    if 'Custom_Create_Date__c' in df.columns:
        date_col_idx = df.columns.get_loc('Custom_Create_Date__c') + 1
        for row in ws.iter_rows(min_row=2, min_col=date_col_idx, max_col=date_col_idx, max_row=ws.max_row):
            for cell in row:
                cell.number_format = 'yyyy-mm-dd'
    
    debug_print("Dates formatted")
    # Right align DATE FLAGGED & DATE CREATED (use renamed names)
    for col_name in ['Date_Flagged__c', 'Custom_Create_Date__c']:
        if col_name in df.columns:
            col_idx = df.columns.get_loc(col_name) + 1
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx, max_row=ws.max_row):
                for cell in row:
                    cell.alignment = Alignment(horizontal='right')

    debug_print("Alignment applied for Date_Flagged__c and Custom_Create_Date__c")
    # Conditional formatting for renamed columns: Survey_Number__c, QuestionID__c, Custom_Create_Date__c
    for col_name in ['Survey_Number__c', 'QuestionID__c', 'Custom_Create_Date__c']:
        if col_name in df.columns:
            col_idx = df.columns.get_loc(col_name) + 1
            col_letter = get_column_letter(col_idx)
            ws.conditional_formatting.add(
                f'{col_letter}2:{col_letter}{ws.max_row}',
                ColorScaleRule(
                    start_type='min', start_color='FF00FF00',
                    mid_type='percentile', mid_value=50, mid_color='CC277BF5',
                    end_type='max', end_color='FFFFFF00'
                )
            )
    
    debug_print("Conditional formatting added for Survey_Number__c, QuestionID__c, Custom_Create_Date__c")
    # Ensure proper data format for sorting - convert all cells to proper data types
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.value is not None:
                # Ensure text values are properly formatted as text
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # string type
                # Ensure numeric values are properly formatted
                elif isinstance(cell.value, (int, float)):
                    cell.data_type = 'n'  # numeric type
    
    debug_print("Ensured proper data format for sorting - convert all cells to proper data types")
    # Update formula logic to use renamed columns
    feedback_col_idx = None
    base_col_idx = None
    if 'Required_Action__c' in df.columns:
        feedback_col_idx = df.columns.get_loc('Required_Action__c') + 1
    if 'QuestionID__c' in df.columns:
        base_col_idx = df.columns.get_loc('QuestionID__c') + 1
    if feedback_col_idx and base_col_idx:
        feedback_col_letter = get_column_letter(feedback_col_idx)
        base_col_letter = get_column_letter(base_col_idx)
        formula = f'ROUND((COUNTA({feedback_col_letter}:{feedback_col_letter})/COUNTA({base_col_letter}:{base_col_letter}))*100,2)&"%"'
        cell = ws.cell(row=1, column=ws.max_column + 1)
        cell.value = f'={formula}'
        if hasattr(cell, 'data_type'):
            cell.data_type = 'f'
    else:
        ws.cell(row=1, column=ws.max_column + 1).value = ''
    
    debug_print("Formula added to last column header to indicate Live Review Progress")
    # Freeze the top row and enable autofilter
    ws.freeze_panes = 'A2'
    # Clear existing autofilter and reapply to ensure clean state
    ws.auto_filter.ref = None
    ws.auto_filter.ref = ws.dimensions
    
    debug_print("Freezed first row")
    debug_print("Excel AutoFilter Enabled")
    # Add data validation dropdown for Setup_Issue__c column
    if 'Setup_Issue__c' in df.columns:
        setup_col_idx = df.columns.get_loc('Setup_Issue__c') + 1
        setup_col_letter = get_column_letter(setup_col_idx)
        
        # Define dropdown options with --None-- as first option
        dropdown_options = [
            "--None--",
            "Custom Instead of Standard",
            "Narrow Custom without Targeting", 
            "Leading (Yes/No)",
            "Leading (Too Few Options)",
            "Indicates PII",
            "Multi-Phase + Other Issues",
            "Grammatical Error / Inconsistent Phrasing",
            "Untranslated Text",
            "Illogical Order",
            "Multiple Errors",
            "Other Set up Issue"
        ]
        
        # Create a helper sheet with the dropdown values
        helper_sheet = wb.create_sheet("ValidationData")
        for i, option in enumerate(dropdown_options, 1):
            helper_sheet.cell(row=i, column=1, value=option)
        
        # Add Required_Action__c dropdown options to the same helper sheet
        req_action_options = [
            "--None--",
            "Rephrase + Need Additional Options",
            "Custom instead of standard", 
            "Language Translation Required"
        ]
        for i, option in enumerate(req_action_options, 1):
            helper_sheet.cell(row=i, column=2, value=option)
        
        # Create data validation for Setup_Issue__c (strict)
        dv = DataValidation(
            type="list",
            formula1=f"ValidationData!$A$1:$A${len(dropdown_options)}",
            allow_blank=False,
            showErrorMessage=True,
            showInputMessage=True
        )
        dv.error = 'You must select a value from the dropdown list only. Manual typing is not allowed.'
        dv.errorTitle = 'Invalid Entry'
        dv.errorStyle = 'stop'  # This prevents invalid entries
        # dv.prompt = 'Please select from the dropdown list only'
        # dv.promptTitle = 'Setup Issue Options'
        
        # Apply to all data rows in the Setup_Issue__c column
        ws.add_data_validation(dv)
        dv.add(f'{setup_col_letter}2:{setup_col_letter}{ws.max_row}')
        
        # Create data validation for Required_Action__c (allows custom values)
        if 'Required_Action__c' in df.columns:
            req_action_col_idx = df.columns.get_loc('Required_Action__c') + 1
            req_action_col_letter = get_column_letter(req_action_col_idx)
            
            dv_req = DataValidation(
                type="list",
                formula1=f"ValidationData!$B$1:$B${len(req_action_options)}",
                allow_blank=True,
                showErrorMessage=False,  # Allow custom values
                showInputMessage=True
            )
            # dv_req.prompt = 'Select from dropdown or type your own value'
            # dv_req.promptTitle = 'Required Action Options'
            
            # Apply to all data rows in the Required_Action__c column
            ws.add_data_validation(dv_req)
            dv_req.add(f'{req_action_col_letter}2:{req_action_col_letter}{ws.max_row}')
        
        # Hide the helper sheet
        helper_sheet.sheet_state = 'hidden'
    
    debug_print("Dropdowns added to Required_Action__c column (typing allowed)")
    debug_print("Dropdowns added to Setup_Issue__c column (strict selection only)")
    # Save to memory buffer and return bytes
    final_buffer = io.BytesIO()
    wb.save(final_buffer)
    final_buffer.seek(0)
    debug_print("Final buffer saved")
    return final_buffer.getvalue()